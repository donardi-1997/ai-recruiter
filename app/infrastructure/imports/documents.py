"""Deterministic CV parsing and safe ZIP expansion for candidate imports."""

from __future__ import annotations

import io
import re
import stat
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Iterator

from app.domains.candidate_imports.rules import (
    document_sha256,
    normalize_email,
    normalize_phone,
)


MIB = 1024 * 1024
MAX_DOCUMENT_BYTES = 15 * MIB
MAX_ARCHIVE_BYTES = 500 * MIB
PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_HEADER_CONTACT_LINE_LIMIT = 40
_HEADER_CHAR_LIMIT = 4000
_NAME_LINE_LIMIT = 8
_NESTED_ARCHIVE_SUFFIXES = {".zip", ".tar", ".gz", ".rar", ".7z"}
_EMAIL_RE = re.compile(r"(?<![\w.+-])([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})(?![\w.-])", re.I)
_PHONE_RE = re.compile(r"(?<!\w)(\+?\d[\d\s().-]{6,}\d)(?!\w)")
_WINDOWS_ABSOLUTE_RE = re.compile(r"^[A-Za-z]:[\\/]")
_GENERIC_NAME_LINES = {
    "curriculum vitae",
    "curriculum",
    "cv",
    "hoja de vida",
    "resume",
    "résumé",
}


class DocumentImportError(ValueError):
    """Base class for deterministic document import failures."""


class UnsupportedDocument(DocumentImportError):
    pass


class InvalidDocument(DocumentImportError):
    pass


class DocumentTooLarge(DocumentImportError):
    pass


class ArchiveTooLarge(DocumentImportError):
    pass


class UnsafeArchive(DocumentImportError):
    pass


class InvalidArchive(DocumentImportError):
    pass


class DocumentLimitExceeded(DocumentImportError):
    pass


class ExpandedArchiveTooLarge(DocumentImportError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    content_type: str
    display_name: str
    email: str | None
    phone: str | None
    text: str
    header_text: str
    sha256: str


@dataclass(frozen=True)
class ExpandedDocument:
    filename: str
    content_type: str
    data: bytes


def _extension(filename: str) -> str:
    lower = filename.casefold()
    if lower.endswith(".pdf"):
        return ".pdf"
    if lower.endswith(".docx"):
        return ".docx"
    if lower.endswith(".zip"):
        return ".zip"
    return ""


def _content_type(filename: str) -> str:
    extension = _extension(filename)
    if extension == ".pdf":
        return PDF_CONTENT_TYPE
    if extension == ".docx":
        return DOCX_CONTENT_TYPE
    raise UnsupportedDocument(f"UNSUPPORTED_DOCUMENT: {filename}")


def _basename(filename: str) -> str:
    normalized = filename.replace("\\", "/")
    return normalized.rsplit("/", 1)[-1]


def _fallback_display_name(filename: str) -> str:
    base = _basename(filename)
    extension = _extension(base)
    stem = base[: -len(extension)] if extension else base
    value = re.sub(r"[-_]+", " ", stem)
    return re.sub(r"\s+", " ", value).strip() or "Unknown"


def _extract_pdf_text(data: bytes) -> str:
    import fitz

    try:
        document = fitz.open(stream=data, filetype="pdf")
        try:
            return "\n".join(page.get_text("text") for page in document).strip()
        finally:
            document.close()
    except Exception as exc:  # library-specific parse errors vary by version
        raise InvalidDocument("INVALID_PDF") from exc


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    except Exception as exc:  # python-docx wraps multiple ZIP/XML failures
        raise InvalidDocument("INVALID_DOCX") from exc


def _build_header(text: str) -> tuple[str, list[str]]:
    selected: list[str] = []
    used_chars = 0
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if len(selected) >= _HEADER_CONTACT_LINE_LIMIT or used_chars >= _HEADER_CHAR_LIMIT:
            break

        separator = 1 if selected else 0
        remaining = _HEADER_CHAR_LIMIT - used_chars - separator
        if remaining <= 0:
            break
        selected.append(line[:remaining])
        used_chars += len(selected[-1]) + separator

    header_text = "\n".join(selected)
    return header_text, selected


def _unique_header_email(lines: list[str]) -> str | None:
    matches = {
        normalize_email(match.group(1))
        for line in lines
        for match in _EMAIL_RE.finditer(line)
    }
    return next(iter(matches)) if len(matches) == 1 else None


def _unique_header_phone(lines: list[str]) -> str | None:
    matches: set[str] = set()
    for line in lines:
        for match in _PHONE_RE.finditer(line):
            normalized = normalize_phone(match.group(1))
            digit_count = sum(character.isdigit() for character in normalized)
            if 7 <= digit_count <= 15:
                matches.add(normalized)
    return next(iter(matches)) if len(matches) == 1 else None


def _plausible_name(line: str) -> bool:
    normalized = re.sub(r"\s+", " ", line).strip()
    if not normalized or normalized.casefold() in _GENERIC_NAME_LINES:
        return False
    if _EMAIL_RE.search(normalized) or _PHONE_RE.search(normalized):
        return False
    if any(character.isdigit() for character in normalized):
        return False
    words = normalized.split()
    if not 2 <= len(words) <= 6:
        return False
    letters = sum(character.isalpha() for character in normalized)
    return letters >= 4 and letters / max(len(normalized), 1) >= 0.6


def _display_name(lines: list[str], filename: str) -> str:
    for line in lines[:_NAME_LINE_LIMIT]:
        if _plausible_name(line):
            return re.sub(r"\s+", " ", line).strip()
    return _fallback_display_name(filename)


def extract_document(data: bytes, filename: str) -> ParsedDocument:
    """Extract deterministic text/header identities from a PDF or DOCX."""
    if not data:
        raise InvalidDocument("EMPTY_DOCUMENT")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise DocumentTooLarge("DOCUMENT_TOO_LARGE")

    content_type = _content_type(filename)
    if content_type == PDF_CONTENT_TYPE:
        text = _extract_pdf_text(data)
    else:
        text = _extract_docx_text(data)

    header_text, header_lines = _build_header(text)
    return ParsedDocument(
        filename=_basename(filename),
        content_type=content_type,
        display_name=_display_name(header_lines, filename),
        email=_unique_header_email(header_lines),
        phone=_unique_header_phone(header_lines),
        text=text,
        header_text=header_text,
        sha256=document_sha256(data),
    )


def _unsafe_archive_path(name: str) -> bool:
    if not name:
        return True
    normalized = name.replace("\\", "/")
    if normalized.startswith("/") or _WINDOWS_ABSOLUTE_RE.match(name):
        return True
    path = PurePosixPath(normalized)
    return path.is_absolute() or any(part == ".." for part in path.parts)


def _is_symlink(info: zipfile.ZipInfo) -> bool:
    if info.create_system != 3:
        return False
    mode = info.external_attr >> 16
    return stat.S_IFMT(mode) == stat.S_IFLNK


def _ignore_archive_entry(name: str) -> bool:
    normalized = name.replace("\\", "/")
    basename = normalized.rsplit("/", 1)[-1]
    return (
        normalized.startswith("__MACOSX/")
        or basename == ".DS_Store"
        or basename.startswith("._")
    )


def _nested_archive(name: str) -> bool:
    suffix = PurePosixPath(name.replace("\\", "/")).suffix.casefold()
    return suffix in _NESTED_ARCHIVE_SUFFIXES


def iter_zip_documents(
    fileobj,
    *,
    remaining_documents: int,
    remaining_bytes: int,
) -> Iterator[ExpandedDocument]:
    """Yield supported CVs from a seekable ZIP one at a time with safety budgets."""
    if remaining_documents < 0 or remaining_bytes < 0:
        raise ValueError("remaining import budgets must be non-negative")

    try:
        archive = zipfile.ZipFile(fileobj, "r")
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise InvalidArchive("INVALID_ARCHIVE") from exc

    expanded_documents = 0
    expanded_bytes = 0

    try:
        for info in archive.infolist():
            name = info.filename
            if _unsafe_archive_path(name):
                raise UnsafeArchive("UNSAFE_ARCHIVE_PATH")
            if _is_symlink(info):
                raise UnsafeArchive("ARCHIVE_SYMLINK")
            if info.is_dir() or _ignore_archive_entry(name):
                continue
            if _nested_archive(name):
                raise UnsafeArchive("NESTED_ARCHIVE")

            extension = _extension(name)
            if extension not in {".pdf", ".docx"}:
                continue

            if expanded_documents + 1 > remaining_documents:
                raise DocumentLimitExceeded("DOCUMENT_LIMIT_EXCEEDED")
            if info.file_size > MAX_DOCUMENT_BYTES:
                raise DocumentTooLarge("DOCUMENT_TOO_LARGE")
            if expanded_bytes + info.file_size > remaining_bytes:
                raise ExpandedArchiveTooLarge("EXPANDED_ARCHIVE_TOO_LARGE")
            if info.flag_bits & 0x1:
                raise UnsafeArchive("ENCRYPTED_ARCHIVE_ENTRY")

            try:
                payload = archive.read(info)
            except (RuntimeError, zipfile.BadZipFile, EOFError) as exc:
                raise InvalidArchive("INVALID_ARCHIVE") from exc

            if len(payload) != info.file_size:
                raise InvalidArchive("ARCHIVE_SIZE_MISMATCH")

            expanded_documents += 1
            expanded_bytes += len(payload)
            yield ExpandedDocument(
                filename=_basename(name),
                content_type=_content_type(name),
                data=payload,
            )
            # The consumer has finished one child before requesting the next.
            # Drop this generator-side reference before reading another payload.
            del payload
    finally:
        archive.close()


def expand_zip(
    data: bytes,
    *,
    remaining_documents: int,
    remaining_bytes: int,
) -> list[ExpandedDocument]:
    """Compatibility wrapper that materializes ZIP children for legacy callers/tests."""
    if len(data) > MAX_ARCHIVE_BYTES:
        raise ArchiveTooLarge("ARCHIVE_TOO_LARGE")
    return list(
        iter_zip_documents(
            io.BytesIO(data),
            remaining_documents=remaining_documents,
            remaining_bytes=remaining_bytes,
        )
    )
